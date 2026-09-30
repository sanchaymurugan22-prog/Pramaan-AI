"""PowerPoint (.pptx) files with python-pptx: a title slide, one slide per generated slide, and a
closing slide. Every slide has speaker notes, the tricolour strip, navy titles, the footer
"AI-assisted · pending human approval" and (if set) the TLP label. The title and closing slides
have an empty box for the QR code (added in Stage 7). Opens in PowerPoint, Keynote and LibreOffice.
"""

from functools import cache
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.exporters.common import FOOTER, PALETTE, QR_PLACEHOLDER, TLP_TEXT_COLOURS, ExportInfo, texts
from app.exporters.fonts import family, font_file

WIDTH, HEIGHT = Inches(13.333), Inches(7.5)  # 16:9
LEFT = Inches(0.8)
CONTENT_WIDTH = WIDTH - 2 * LEFT

FOOTER_SIZE = 11             # points
FOOTER_BOX = Inches(7)       # the right-hand footer box ("Job #12 · title · date · 3/6")
# Room for text inside that box: its width minus the 0.1-inch inner margins, less 10% in case the
# computer that opens the deck does not have Hind installed and uses a slightly wider font.
FOOTER_TEXT_POINTS = (FOOTER_BOX / 12700 - 2 * 7.2) * 0.9


def write_pptx(info: ExportInfo, content: dict, path: Path) -> Path:
    deck = Presentation()
    deck.slide_width, deck.slide_height = WIDTH, HEIGHT
    title = content.get("title") or info.job_title
    slides = content.get("slides") or []
    total = len(slides) + 2

    _title_slide(deck, info, title, len(slides), total)
    for number, slide in enumerate(slides, start=2):
        _content_slide(deck, info, slide, number, total)
    _closing_slide(deck, info, title, total)

    _list_notes_master(deck)
    props = deck.core_properties
    props.title, props.subject, props.author = title, info.job_title, "Pramaan AI"
    props.comments = f"{FOOTER}. {info.header_line()}"
    deck.save(path)
    return path


# ---- slides -------------------------------------------------------------------------------


def _new_slide(deck, info: ExportInfo, number: int, total: int, notes: str):
    """A blank slide with the parts every slide has: strip, TLP label, footer, speaker notes."""
    slide = deck.slides.add_slide(deck.slide_layouts[6])  # layout 6 = blank
    third = WIDTH // 3
    for index, colour in enumerate(("saffron", "white", "green")):
        _rect(slide, index * third, 0, third if index < 2 else WIDTH - 2 * third, Inches(0.1), colour)
    if info.tlp_label:
        label = _rect(slide, WIDTH - LEFT - Inches(1.5), Inches(0.3), Inches(1.5), Inches(0.36), "black")
        _write(label.text_frame, info.tlp_label, "heading", 12, info, colour=TLP_TEXT_COLOURS.get(info.tlp, "FFFFFF"),
               bold=True, align=PP_ALIGN.CENTER)
        label.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    _rect(slide, LEFT, HEIGHT - Inches(0.62), CONTENT_WIDTH, Pt(1), "line")
    _text(slide, LEFT, HEIGHT - Inches(0.55), Inches(5), Inches(0.35), FOOTER, "heading", 11, info,
          colour="saffron_dark", bold=True)
    box = _text(slide, WIDTH - LEFT - FOOTER_BOX, HEIGHT - Inches(0.55), FOOTER_BOX, Inches(0.35),
                footer_text(info, number, total), "body", FOOTER_SIZE, info, colour="muted", align=PP_ALIGN.RIGHT)
    box.text_frame.word_wrap = False  # one line, always (footer_text already made it fit)

    slide.notes_slide.notes_text_frame.text = notes
    return slide


def footer_text(info: ExportInfo, number: int, total: int) -> str:
    """'Job #12 · Hospital ransomware · 30 Sep 2026 · 3/6'. A long job title is shortened with '…' so the
    whole footer fits on one line; the date and slide number are always shown in full."""
    start, end = f"Job #{info.job_id} · ", f" · {info.date} · {number}/{total}"
    title = info.job_title.strip()
    text = f"{start}{title}{end}"
    while title and footer_width(text, info.language) > FOOTER_TEXT_POINTS:
        # drop the last word (or the last letter of a single very long word)
        title = (title.rsplit(" ", 1)[0] if " " in title else title[:-1]).rstrip(" ·,:;-")
        text = f"{start}{title}…{end}"
    return text


def footer_width(text: str, language: str = "en") -> float:
    """Width of footer text in points, measured with the bundled body font (Hind)."""
    return _body_font(language).getlength(text) * FOOTER_SIZE / 1000


@cache
def _body_font(language: str):
    return ImageFont.truetype(str(font_file("body", "regular", language)), size=1000)


def _title_slide(deck, info: ExportInfo, title: str, count: int, total: int) -> None:
    notes = (f"Introduce the briefing: {title}. It is based on '{info.job_title}' and has {count} content slides. "
             f"This deck is AI-assisted and pending human approval until a reviewer signs it.")
    slide = _new_slide(deck, info, 1, total, notes)
    _text(slide, LEFT, Inches(1.9), Inches(9), Inches(0.4), f"PRESENTATION · JOB #{info.job_id}", "heading", 14, info,
          colour="saffron_dark", bold=True)
    _text(slide, LEFT, Inches(2.4), Inches(9.2), Inches(2.2), title, "heading", 40 if len(title) < 50 else 32, info,
          colour="navy", bold=True, anchor=MSO_ANCHOR.TOP)
    _rect(slide, LEFT, Inches(4.75), Inches(1.2), Inches(0.08), "saffron")
    _rect(slide, LEFT + Inches(1.2), Inches(4.75), Inches(1.2), Inches(0.08), "green")
    _text(slide, LEFT, Inches(5.0), Inches(9), Inches(0.9), f"{info.job_title}\nPrepared {info.date}", "body", 18, info,
          colour="muted")
    _qr_placeholder(slide, info, WIDTH - LEFT - Inches(2.1), Inches(2.5), Inches(2.1))


def _content_slide(deck, info: ExportInfo, content: dict, number: int, total: int) -> None:
    notes = content.get("speaker_notes") or "No speaker notes were written for this slide."
    slide = _new_slide(deck, info, number, total, notes)
    title = content.get("title", "")
    _text(slide, LEFT, Inches(0.75), CONTENT_WIDTH - Inches(1.8), Inches(1.2), title, "heading",
          32 if len(title) < 45 else 26, info, colour="navy", bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, LEFT, Inches(2.05), Inches(1.0), Inches(0.07), "saffron")

    bullets = texts(content.get("bullets"))
    longest = max((len(b) for b in bullets), default=0)
    size = 24 if len(bullets) <= 3 and longest < 70 else 20
    box = slide.shapes.add_textbox(LEFT, Inches(2.45), CONTENT_WIDTH, Inches(4.2))
    frame = box.text_frame
    frame.word_wrap = True
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(14)
        _add_run(paragraph, "●   ", "heading", size - 6, info, colour="saffron")
        _add_run(paragraph, bullet, "body", size, info)


def _closing_slide(deck, info: ExportInfo, title: str, total: int) -> None:
    notes = ("Thank the audience and invite questions. Remind them that this deck is AI-assisted and pending human "
             "approval: once signed, the QR code on this slide lets anyone check it is genuine.")
    slide = _new_slide(deck, info, total, total, notes)
    _text(slide, LEFT, Inches(2.2), Inches(9), Inches(1.2), "Thank you", "heading", 48, info, colour="navy", bold=True)
    _text(slide, LEFT, Inches(3.4), Inches(9), Inches(0.6), "Questions and discussion", "heading", 24, info,
          colour="saffron_dark")
    _text(slide, LEFT, Inches(4.3), Inches(9), Inches(1.2),
          f"{title}\nScan the QR code to check this deck is genuine (added when signed).", "body", 16, info, colour="muted")
    _qr_placeholder(slide, info, WIDTH - LEFT - Inches(2.1), Inches(2.3), Inches(2.1))


def _list_notes_master(deck) -> None:
    """python-pptx links the speaker-notes master but does not list it in presentation.xml, as
    PowerPoint does. Without the list, macOS Quick Look (spacebar preview) hangs on the file."""
    presentation = deck.part._element
    if presentation.find(qn("p:notesMasterIdLst")) is not None:
        return
    rel_id = next((r.rId for r in deck.part.rels.values() if r.reltype == RT.NOTES_MASTER), None)
    if rel_id is None:
        return
    id_list = presentation.makeelement(qn("p:notesMasterIdLst"), {})
    id_list.append(id_list.makeelement(qn("p:notesMasterId"), {qn("r:id"): rel_id}))
    presentation.find(qn("p:sldMasterIdLst")).addnext(id_list)  # it must come right after the slide masters


# ---- drawing helpers ----------------------------------------------------------------------


def _qr_placeholder(slide, info: ExportInfo, left, top, size) -> None:
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, size, size)
    box.shadow.inherit = False
    box.fill.background()
    box.line.color.rgb = _colour("line_2")
    box.line.width = Pt(1.5)
    box.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    frame = box.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    first, second = QR_PLACEHOLDER.split("\n")
    _write(frame, first, "heading", 14, info, colour="muted", bold=True, align=PP_ALIGN.CENTER)
    paragraph = frame.add_paragraph()
    paragraph.alignment = PP_ALIGN.CENTER
    _add_run(paragraph, second, "body", 12, info, colour="muted")


def _colour(name: str) -> RGBColor:
    return RGBColor.from_string(PALETTE.get(name, name))


def _rect(slide, left, top, width, height, colour: str):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    shape.fill.solid()
    shape.fill.fore_color.rgb = _colour(colour)
    shape.line.fill.background()
    return shape


def _text(slide, left, top, width, height, text: str, role: str, size: float, info: ExportInfo,
          colour: str = "ink", bold: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    for index, line in enumerate(text.split("\n")):
        if index == 0:
            _write(frame, line, role, size, info, colour=colour, bold=bold, align=align)
        else:
            paragraph = frame.add_paragraph()
            paragraph.alignment = align
            _add_run(paragraph, line, role, size, info, colour=colour, bold=bold)
    return box


def _write(frame, text: str, role: str, size: float, info: ExportInfo, colour: str = "ink", bold: bool = False,
           align=PP_ALIGN.LEFT) -> None:
    """Put text in the first paragraph of a text frame."""
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    _add_run(paragraph, text, role, size, info, colour=colour, bold=bold)


def _add_run(paragraph, text: str, role: str, size: float, info: ExportInfo, colour: str = "ink", bold: bool = False):
    run = paragraph.add_run()
    run.text = text
    font = run.font
    font.name = family(role, info.language)
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = _colour(colour)
    # The same font for Indian scripts ("cs" = complex script), so Stage 8 languages use it too.
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        element = rpr.find(qn(tag))
        if element is None:
            element = rpr.makeelement(qn(tag), {})
            rpr.insert_element_before(element, "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")
        element.set("typeface", font.name)
    return run
