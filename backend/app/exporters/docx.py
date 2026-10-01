"""Word (.docx) files with python-docx: advisory, executive summary, and the video script/storyboard.

Page layout: tricolour strip and TLP label in the page header; a top block with the title, the
job details and an empty box for the QR code (added in Stage 7); then the sections from
blocks.py; the footer "AI-assisted · pending human approval" with the job and page number.
"""

import io
from pathlib import Path

import docx
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from app.exporters.blocks import Block, document_blocks
from app.exporters.common import PALETTE, QR_PLACEHOLDER, TLP_TEXT_COLOURS, ExportInfo, fit_title, rgb
from app.exporters.fonts import family, text_width

PAGE_WIDTH, PAGE_HEIGHT, MARGIN = Cm(21), Cm(29.7), Cm(2)  # A4
TEXT_WIDTH = PAGE_WIDTH - 2 * MARGIN
QR_BOX = Cm(3.2)


def write_docx(info: ExportInfo, content: dict, path: Path) -> Path:
    blocks = document_blocks(info.output_type, content)
    title = next((b.text for b in blocks if b.kind == "title"), info.output_label)
    subtitle = next((b.text for b in blocks if b.kind == "subtitle"), "")

    document = docx.Document()
    _page_setup(document, info)
    _top_block(document, info, title, subtitle)
    for block in blocks:
        if block.kind not in ("title", "subtitle"):
            _add_block(document, block, info)

    props = document.core_properties
    props.title, props.subject, props.author = title, info.job_title, "Pramaan AI"
    props.comments = f"{info.footer}. {info.header_line()}"
    document.save(path)
    return path


# ---- page, header, footer -----------------------------------------------------------------


def _page_setup(document, info: ExportInfo) -> None:
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.font.color.rgb = _colour("ink")
    _set_font_name(normal.element.get_or_add_rPr(), family("body", info.language))
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    section = document.sections[0]
    section.page_width, section.page_height = PAGE_WIDTH, PAGE_HEIGHT
    section.left_margin = section.right_margin = section.bottom_margin = MARGIN
    section.top_margin = Cm(2.2)
    section.header_distance, section.footer_distance = Cm(0.6), Cm(1)

    # Word's tab stops every 0.7 cm (lines up the bullets in every app, not just Word)
    default_tab = document.settings.element.find(qn("w:defaultTabStop"))
    if default_tab is not None:
        default_tab.set(qn("w:val"), str(int(Cm(0.7).twips)))

    # Header: the tricolour strip (a small picture, so every app shows it), then the TLP label.
    header = section.header
    strip = header.paragraphs[0]
    strip.paragraph_format.space_after = Pt(0)
    strip.add_run().add_picture(_strip_image(), width=TEXT_WIDTH, height=Pt(3))

    line = header.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    line.paragraph_format.space_before = Pt(4)
    _run(line, "Pramaan AI · " + info.output_label, "heading", 8, info, colour="muted", bold=True)
    if info.tlp_label:
        line.add_run("   ")
        _tlp_run(line, info)

    footer = section.footer.paragraphs[0]
    _run(footer, info.footer, "heading", 8.5, info, colour="saffron_dark", bold=True)
    # One line: a long job title is shortened with "…" (room for a 3-digit page number is kept)
    room = TEXT_WIDTH / 12700 - text_width(info.footer, 8.5, "heading", "bold", info.language) - 6
    job_line = fit_title(f"  ·  Job #{info.job_id}: ", info.job_title, "  ·  Page 999",
                         lambda text: text_width(text, 8.5, "body", "regular", info.language), room)
    _run(footer, job_line.removesuffix("999"), "body", 8.5, info, colour="muted")
    _page_number(footer)


def _top_block(document, info: ExportInfo, title: str, subtitle: str) -> None:
    """Title and job details on the left, the empty QR code box on the right."""
    table = document.add_table(rows=1, cols=2)
    _column_widths(table, [TEXT_WIDTH - QR_BOX - Cm(0.4), QR_BOX + Cm(0.4)])
    left, right = table.rows[0].cells
    table.rows[0].height, table.rows[0].height_rule = QR_BOX, WD_ROW_HEIGHT_RULE.AT_LEAST

    eyebrow = left.paragraphs[0]
    _run(eyebrow, f"{info.output_label.upper()} · JOB #{info.job_id}", "heading", 9, info, colour="saffron_dark", bold=True)
    eyebrow.paragraph_format.space_after = Pt(2)
    heading = left.add_paragraph()
    _run(heading, title, "heading", 20, info, colour="navy", bold=True)
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.line_spacing = 1.0
    if subtitle:
        _run(left.add_paragraph(), subtitle, "body", 11, info, colour="red" if "HIGH" in subtitle or "CRITICAL" in subtitle else "muted", bold=True)
    meta = left.add_paragraph()
    _run(meta, f"{info.job_title} · Prepared {info.date}", "body", 10, info, colour="muted")
    if info.tlp_label:
        meta.add_run("   ")
        _tlp_run(meta, info)

    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if info.signed:  # Stage 7: the real QR code (verify address + record number) and the record number
        picture = right.paragraphs[0]
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_after = Pt(0)
        picture.add_run().add_picture(io.BytesIO(info.qr_png(8)), width=QR_BOX - Cm(0.2))
        caption = right.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(caption, info.record_no, "mono", 7, info, colour="muted")
        document.add_paragraph().paragraph_format.space_after = Pt(2)
        return
    # QR code placeholder until signed: a dashed box.
    _borders(right, "dashed", PALETTE["line_2"], size=12)
    for number, line in enumerate(QR_PLACEHOLDER.split("\n")):
        paragraph = right.paragraphs[0] if number == 0 else right.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        _run(paragraph, line, "heading" if number == 0 else "body", 9 if number == 0 else 8, info, colour="muted", bold=number == 0)

    document.add_paragraph().paragraph_format.space_after = Pt(2)


# ---- blocks -------------------------------------------------------------------------------


def _add_block(document, block: Block, info: ExportInfo) -> None:
    if block.kind == "heading":
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
        _run(paragraph, block.text, "heading", 13, info, colour="navy", bold=True)
    elif block.kind == "paragraph":
        _run(document.add_paragraph(), block.text, "body", 11, info)
    elif block.kind in ("bullets", "numbered"):
        for number, item in enumerate(block.items, start=1):
            paragraph = document.add_paragraph()
            fmt = paragraph.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Cm(0.7), Cm(-0.7)
            fmt.tab_stops.add_tab_stop(Cm(0.7))
            fmt.space_after = Pt(4)
            marker = "•" if block.kind == "bullets" else f"{number}."
            _run(paragraph, marker + "\t", "heading", 11, info, colour="saffron_dark", bold=True)
            _run(paragraph, item, "body", 11, info)
    elif block.kind == "callout":
        table = document.add_table(rows=1, cols=1)
        _column_widths(table, [TEXT_WIDTH])
        cell = table.rows[0].cells[0]
        _shade(cell, PALETTE["saffron_light"])
        _borders(cell, "none", "auto")
        _border(cell, "left", "single", PALETTE["saffron"], size=24)
        label = cell.paragraphs[0]
        _run(label, block.label.upper(), "heading", 9, info, colour="saffron_dark", bold=True)
        label.paragraph_format.space_after = Pt(2)
        _run(cell.add_paragraph(), block.text, "body", 12, info)
        document.add_paragraph().paragraph_format.space_after = Pt(0)
    elif block.kind == "table":
        _data_table(document, block, info)


def _data_table(document, block: Block, info: ExportInfo) -> None:
    table = document.add_table(rows=1 + len(block.rows), cols=len(block.header))
    table.style = document.styles["Table Grid"]
    _table_borders(table, PALETTE["line_2"])
    widths = [int(TEXT_WIDTH * share) for share in (block.widths or [1 / len(block.header)] * len(block.header))]
    _column_widths(table, widths)
    small = len(block.header) > 3  # the storyboard has 5 columns: use smaller text
    for row_number, values in enumerate([block.header, *block.rows]):
        for column, (cell, value) in enumerate(zip(table.rows[row_number].cells, values)):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            if row_number == 0:
                _shade(cell, PALETTE["navy_light"])
                _run(paragraph, value, "heading", 9, info, colour="navy", bold=True)
            else:
                role = "mono" if column in block.mono_columns else "body"
                _run(paragraph, value, role, 9 if role == "mono" or small else 10.5, info)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


# ---- small python-docx helpers --------------------------------------------------------------


def _strip_image() -> io.BytesIO:
    """A 300 x 4 pixel PNG: saffron | white | green."""
    image = Image.new("RGB", (300, 4), rgb("white"))
    image.paste(rgb("saffron"), (0, 0, 100, 4))
    image.paste(rgb("green"), (200, 0, 300, 4))
    stream = io.BytesIO()
    image.save(stream, "PNG")
    stream.seek(0)
    return stream


def _column_widths(table, widths: list[int]) -> None:
    """Fixed column widths. Word reads each cell's width; Pages and Quick Look read the table's
    column list ("grid"), so both are set."""
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = width
        for cell in column.cells:
            cell.width = width


def _colour(name: str) -> RGBColor:
    return RGBColor.from_string(PALETTE.get(name, name))


def _set_font_name(rpr, name: str) -> None:
    """Set the font for Latin text and for Indian scripts ("cs" = complex script) alike."""
    fonts = rpr.get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attribute), name)


def _run(paragraph, text: str, role: str, size: float, info: ExportInfo, colour: str = "ink", bold: bool = False):
    run = paragraph.add_run(text)
    _set_font_name(run._element.get_or_add_rPr(), family(role, info.language))
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _colour(colour)
    return run


def _tlp_run(paragraph, info: ExportInfo) -> None:
    """The TLP label: coloured letters on black, as the TLP standard asks."""
    run = _run(paragraph, f" {info.tlp_label} ", "heading", 9, info, colour=TLP_TEXT_COLOURS.get(info.tlp, "FFFFFF"), bold=True)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), PALETTE["black"])
    run._element.get_or_add_rPr().append(shading)


def _page_number(paragraph) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


# Word is strict about the order of settings inside a cell or table, so new settings are put
# before the ones that must come after them.
_AFTER_CELL_SHADING = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
_AFTER_TABLE_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")


def _shade(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().insert_element_before(shading, *_AFTER_CELL_SHADING)


def _border(cell, side: str, style: str, colour: str, size: int = 8) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.insert_element_before(borders, "w:shd", *_AFTER_CELL_SHADING)
    edge = borders.find(qn(f"w:{side}"))
    if edge is None:  # _borders() adds the sides in Word's order: top, left, bottom, right
        edge = OxmlElement(f"w:{side}")
        borders.append(edge)
    edge.set(qn("w:val"), style)
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:color"), colour)


def _borders(cell, style: str, colour: str, size: int = 8) -> None:
    for side in ("top", "left", "bottom", "right"):
        _border(cell, side, style, colour, size)


def _table_borders(table, colour: str) -> None:
    """Thin borders in the design's line colour instead of Word's black grid."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), colour)
        borders.append(edge)
    tbl_pr.insert_element_before(borders, *_AFTER_TABLE_BORDERS)
