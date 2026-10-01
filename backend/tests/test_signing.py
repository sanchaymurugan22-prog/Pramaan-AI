"""Stage 7 part 1: signing on Approve. Record number, final files with a QR code, file and text
fingerprints, signed manifest, read-only signed files, the signed kit, and a new version."""

import hashlib
import io
import json
import zipfile

import docx
import pytest

from app import crypto
from app.config import settings
from app.db import Record, SessionLocal
from app.signing import records
from app.signing.sign_job import signed_dir
from app.signing.signer import TestSigner, get_signer, verify
from app.signing.texts import normalise, text_hash
from tests.auth_helpers import signed_in_client
from tests.helpers import SAMPLE_REPORT
from tests.test_jobs_api import wait_for

operator = signed_in_client("operator", "sign.operator")
reviewer = signed_in_client("reviewer", "sign.reviewer")


def approved_job(outputs=("x_thread", "advisory"), tlp: str | None = None, title: str = "") -> dict:
    """A job taken through submit and approve (which signs it)."""
    created = operator.post("/api/jobs", data={"text": SAMPLE_REPORT.read_text(encoding="utf-8"), "title": title, **(
        {"outputs": list(outputs)} if tlp is None else {})})
    job_id = created.json()["id"]
    if tlp is not None:  # the Safety check, with this label
        assert operator.put(f"/api/jobs/{job_id}/safety", json={"tlp": tlp}).status_code == 200
        assert operator.post(f"/api/jobs/{job_id}/start", json={"outputs": list(outputs)}).status_code == 200
    wait_for(job_id)
    assert operator.post(f"/api/jobs/{job_id}/submit", json={}).status_code == 200
    response = reviewer.post(f"/api/jobs/{job_id}/review", json={"decision": "approve", "notes": "Fine"})
    assert response.status_code == 200, response.text
    return response.json()


def entry_for(record_no: str) -> Record:
    with SessionLocal() as db:
        return records.find_issue(db, record_no)


def test_approve_signs_the_job():
    job = approved_job()
    record = job["record"]
    assert job["status"] == "approved"
    assert record["record_no"].startswith("PRM-") and len(record["record_no"]) == len("PRM-2026-000001")
    assert record["approved_by"] == "Sign Reviewer" and record["current"]
    assert record["verify_url"] == f"{settings.verify_base_url}/?r={record['record_no']}"
    # x_thread -> txt; advisory -> pdf + docx
    assert sorted(f["name"].rsplit(".", 1)[1] for f in record["files"]) == ["docx", "pdf", "txt"]

    entry = entry_for(record["record_no"])
    manifest = json.loads(entry.manifest)
    assert manifest["title"] == job["title"] and manifest["approved_by"]["role"] == "Reviewer"
    assert manifest["issuing_office"] == settings.issuing_office and manifest["job_version"] == 1
    assert {t["label"] for t in manifest["texts"]} >= {"Whole thread", "Post 1", "Full text"}


def test_the_signature_verifies_and_one_changed_byte_fails():
    job = approved_job(outputs=("x_thread",))
    entry = entry_for(job["record"]["record_no"])
    public_key = get_signer().public_key_pem()
    assert verify(public_key, entry.manifest.encode(), entry.signature)
    assert verify(public_key, entry.public_manifest.encode(), entry.public_signature)

    changed = bytearray(entry.manifest.encode())
    changed[10] ^= 1
    assert not verify(public_key, bytes(changed), entry.signature)
    assert not verify(public_key, entry.manifest.replace("Sign Reviewer", "Someone Else").encode(), entry.signature)
    assert not verify(public_key, entry.manifest.encode(), entry.public_signature)  # signatures don't swap
    assert not verify(public_key, entry.manifest.encode(), "not base64!")


def test_downloads_are_the_signed_files_with_the_qr_code():
    job = approved_job()
    record_no = job["record"]["record_no"]
    hashes = {f["name"]: f["sha256"] for f in job["record"]["files"]}
    for output in job["outputs"]:
        for fmt in output["formats"]:
            response = operator.get(f"/api/jobs/{job['id']}/outputs/{output['id']}/download?format={fmt}")
            assert response.status_code == 200
            name = response.headers["content-disposition"].split('filename="')[1].rstrip('"')
            assert hashlib.sha256(response.content).hexdigest() == hashes[name]  # exactly what was fingerprinted
            if fmt == "txt":
                text = response.content.decode()
                assert f"Approved and signed · Record {record_no}" in text
                assert f"Check it is genuine: {settings.verify_base_url}/?r={record_no}" in text
            if fmt == "docx":
                document = docx.Document(io.BytesIO(response.content))
                assert len(document.inline_shapes) == 1  # the QR code (the tricolour strip is in the page header)
                assert any(record_no in p.text for cell in document.tables[0].rows[0].cells for p in cell.paragraphs)


def test_signed_files_are_stored_encrypted_and_read_only():
    job = approved_job(outputs=("x_thread",))
    folder = signed_dir(job["id"], job["record"]["record_no"])
    files = list(folder.iterdir())
    assert files and all(f.read_bytes().startswith(crypto.MAGIC) for f in files)
    assert all(f.stat().st_mode & 0o222 == 0 for f in files)  # nobody may write
    output = job["outputs"][0]
    edit = {"fields": [{"path": ["tweets", 0, "text"], "text": "Changed after signing"}]}
    assert operator.put(f"/api/jobs/{job['id']}/outputs/{output['id']}", json=edit).status_code == 409


def test_the_signed_kit_can_be_checked_offline():
    job = approved_job(outputs=("x_thread", "executive_summary"))
    kit = zipfile.ZipFile(io.BytesIO(operator.get(f"/api/jobs/{job['id']}/kit.zip").content))
    names = set(kit.namelist())
    assert {"README.txt", "record.json", "public-key.pem"} <= names
    record = json.loads(kit.read("record.json"))
    assert verify(kit.read("public-key.pem").decode(), record["manifest"].encode(), record["signature"])
    published = json.loads(record["manifest"])
    for f in published["files"]:
        assert hashlib.sha256(kit.read(f["name"])).hexdigest() == f["sha256"]
    assert job["record"]["record_no"] in kit.read("README.txt").decode()


def test_tlp_red_and_amber_publish_no_title_or_text():
    for tlp in ("AMBER", "RED"):
        job = approved_job(outputs=("advisory",), tlp=tlp)
        entry = entry_for(job["record"]["record_no"])
        public = json.loads(entry.public_manifest)
        assert public["restricted"] is True
        assert "title" not in public and job["title"] not in entry.public_manifest
        assert set(public) == {"format", "kind", "record_no", "issued_at", "replaces", "signer", "restricted",
                               "files", "texts"}
        assert all(set(f) == {"sha256"} for f in public["files"]) and all(set(t) == {"sha256"} for t in public["texts"])
        assert json.loads(entry.manifest)["title"] == job["title"]  # still in the full, internal record


def test_tlp_green_publishes_title_role_and_message_texts():
    job = approved_job(outputs=("x_thread", "advisory"), tlp="GREEN")
    public = json.loads(entry_for(job["record"]["record_no"]).public_manifest)
    assert public["restricted"] is False and public["title"] == job["title"]
    assert public["approved_by"] == {"role": "Reviewer"}  # the role, not the name
    assert "Sign Reviewer" not in json.dumps(public)
    thread = [t for t in public["texts"] if t["output"] == "x_thread"]
    assert thread and all("text" in t and t["sha256"] == text_hash(t["text"]) for t in thread)
    assert all("text" not in t for t in public["texts"] if t["output"] == "advisory")


def test_a_new_version_needs_a_new_signature_and_replaces_the_old_record():
    job = approved_job(outputs=("x_thread",))
    first = job["record"]["record_no"]
    url = f"/api/jobs/{job['id']}"
    reopened = operator.post(f"{url}/new-version")
    assert reopened.status_code == 200
    assert reopened.json()["status"] == "ready" and reopened.json()["version"] == 2
    assert reopened.json()["record"]["current"] is False
    output = job["outputs"][0]
    edit = {"fields": [{"path": ["tweets", 0, "text"], "text": "Updated: patch your VPN gateway today."}]}
    assert operator.put(f"{url}/outputs/{output['id']}", json=edit).status_code == 200
    # unsigned again until approved: downloads are not the signed files
    text = operator.get(f"{url}/outputs/{output['id']}/download?format=txt").text
    assert "pending human approval" in text

    operator.post(f"{url}/submit", json={})
    second = reviewer.post(f"{url}/review", json={"decision": "approve"}).json()["record"]
    assert second["record_no"] != first and second["version"] == 2
    assert json.loads(entry_for(second["record_no"]).manifest)["replaces"] == first
    assert entry_for(first) is not None  # the old record is kept


def test_the_private_key_is_stored_encrypted(tmp_path):
    signer = TestSigner(tmp_path)
    public = signer.public_key_pem()
    stored = (tmp_path / "test-signer.key").read_bytes()
    assert stored.startswith(crypto.MAGIC) and b"PRIVATE KEY" not in stored
    assert "PUBLIC KEY" in public and verify(public, b"hello", signer.sign(b"hello"))
    assert len(signer.key_id()) == 16


@pytest.mark.parametrize("a, b", [
    ("Patch your VPN today!", "patch   your vpn today"),
    ("Patch your VPN today 🙏🔒", "Patch your VPN today"),
    ("Patch\u200b your VPN\u200d today", "Patch your VPN today"),
    ("ＰＡＴＣＨ your VPN today", "patch your vpn today"),
    ("Call 1930, or visit cybercrime.gov.in.", "call 1930 or visit cybercrime gov in"),
    ("तुरंत पैच करें!", "तुरंत पैच करें"),
])
def test_normalised_text_ignores_spaces_case_punctuation_emojis_and_hidden_characters(a, b):
    assert normalise(a) == normalise(b)
    assert text_hash(a) == text_hash(b)


def test_normalising_keeps_the_words():
    assert normalise("Patch your VPN today") != normalise("Patch your VPN tomorrow")
    assert normalise("तुरंत पैच करें") == "तुरंत पैच करें"  # Hindi vowel signs are kept
