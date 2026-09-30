"""Tests for reading sources: pasted text, .txt, .pdf and .docx, with page numbers."""

import io

import docx
import pytest
from tests.helpers import SAMPLE_REPORT, make_pdf

from app.pipeline import ingest


def test_pasted_text_is_one_page():
    source = ingest.from_text("Hello\nworld")
    assert source.kind == "text"
    assert source.pages == ["Hello\nworld"]
    assert len(source.sha256) == 64


def test_txt_form_feed_starts_a_new_page():
    source = ingest.from_file("report.txt", "Page one text\fPage two text".encode())
    assert source.pages == ["Page one text", "Page two text"]


def test_sample_report_reads():
    source = ingest.from_file(SAMPLE_REPORT.name, SAMPLE_REPORT.read_bytes())
    assert "NightLedger" in source.pages[0]
    assert source.pages[0].startswith("SAMPLE – FICTIONAL")  # the en dash survives decoding


def test_pdf_keeps_pages():
    source = ingest.from_file("report.pdf", make_pdf(["First page about hospitals", "Second page about backups"]))
    assert source.kind == "pdf"
    assert len(source.pages) == 2
    assert "hospitals" in source.pages[0]
    assert "backups" in source.pages[1]


def test_docx_splits_at_page_breaks_and_reads_tables():
    document = docx.Document()
    document.add_paragraph("Intro paragraph")
    document.add_page_break()
    document.add_paragraph("Second page paragraph")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "IP"
    table.rows[0].cells[1].text = "203.0.113.45"
    buffer = io.BytesIO()
    document.save(buffer)

    source = ingest.from_file("report.docx", buffer.getvalue())
    assert source.pages[0] == "Intro paragraph"
    assert "Second page paragraph" in source.pages[1]
    assert "IP | 203.0.113.45" in source.pages[1]


def test_scanned_pdf_gets_a_friendly_message():
    with pytest.raises(ingest.IngestError, match="scanned"):
        ingest.from_file("scan.pdf", make_pdf([""]))


def test_unsupported_file_type():
    with pytest.raises(ingest.IngestError, match="only .txt, .pdf and .docx"):
        ingest.from_file("photo.png", b"\x89PNG")


def test_save_and_load_pages(tmp_path, monkeypatch):
    monkeypatch.setattr(ingest.settings, "data_dir", tmp_path)
    source = ingest.from_text("a\fb")
    path = ingest.save_source(7, "S1", source)
    assert ingest.load_pages(path) == ["a", "b"]
    assert (tmp_path / "jobs" / "7" / "sources" / "S1-pasted-text.txt").read_bytes() == b"a\fb"
