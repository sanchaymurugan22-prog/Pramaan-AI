"""Stage 5 tests (mock AI): source highlights, "Not in source", consistency across outputs, quality
score, human edits with versions, regenerating one output, and the two export fixes."""

import io
import re

import pytest
from fastapi.testclient import TestClient

from tests.auth_helpers import signed_in_client
from pptx import Presentation
from pypdf import PdfReader

from app.exporters.common import ExportInfo
from app.exporters.pptx import FOOTER_TEXT_POINTS, footer_text, footer_width, write_pptx
from app.main import app
from app.pipeline import checks
from app.pipeline.factsheet import SourcePages, build_fact_sheet
from app.pipeline.trace import locate
from tests.canned import canned_ai, reply
from tests.helpers import SAMPLE_REPORT
from tests.test_jobs_api import wait_for

client = signed_in_client("operator")


@pytest.fixture(scope="module", autouse=True)
def fixed_answers():
    """These tests check scores, edits and files against known texts, so the AI gives fixed answers
    (tests/canned.py) instead of building them from the source."""
    with canned_ai():
        yield
SAMPLE = SAMPLE_REPORT.read_text(encoding="utf-8")
SHA256 = "9f2c4b7e1a3d5f60718293a4b5c6d7e8f9012a3b4c5d6e7f8091a2b3c4d5e41a"


@pytest.fixture(scope="module")
def sheet() -> dict:
    return build_fact_sheet([SourcePages("S1", SAMPLE_REPORT.name, [SAMPLE])])


@pytest.fixture(scope="module")
def known(sheet):
    return checks.known_values([SAMPLE])


def words(text: str) -> list[str]:
    return re.findall(r"\w+", text.lower())


# ---- 1. highlight offsets ------------------------------------------------------------------------


def test_exact_quote_gives_the_characters_to_highlight():
    pages = [("S1", 1, "Nothing here."), ("S1", 2, "As of 28\nSeptember, 42 hospitals in five states have reported disruption.")]
    found = locate("42 hospitals in five states have reported disruption", pages)
    assert (found.found, found.source_id, found.page) == ("exact", "S1", 2)
    assert pages[1][2][found.start : found.end] == "42 hospitals in five states have reported disruption"

    # punctuation, capitals and line breaks do not matter
    found = locate("as of 28 September 42 Hospitals", pages)
    assert found.found == "exact"
    assert pages[1][2][found.start : found.end] == "As of 28\nSeptember, 42 hospitals"


def test_close_and_missing_quotes():
    page = "As of 28 September, 42 hospitals in five states have reported disruption, mostly billing."
    close = locate("As of 28 September, 42 hospitals in 5 states have reported disruption", [("S1", 1, page)])
    assert close.found == "close"
    assert page[close.start : close.end].startswith("As of 28 September")
    assert page[close.start : close.end].endswith("disruption")

    missing = locate("Attackers demand payment in cryptocurrency", [("S1", 1, page)])
    assert (missing.found, missing.start, missing.end) == ("no", None, None)


def test_every_fact_action_and_date_of_the_sample_is_highlighted(sheet):
    for key in ("key_facts", "recommended_actions", "dates"):
        for item in sheet[key]:
            assert item["quote_found"] in ("exact", "close"), item
            highlighted = SAMPLE[item["start"] : item["end"]]
            assert highlighted and words(highlighted) == words(item["quote"]), item["id"]
    assert [d["id"] for d in sheet["dates"]][:2] == ["D1", "D2"]


def test_sentences_are_split_without_breaking_numbers():
    text = "About 1.2 million records were copied. Patch now! e.g. this stays. New line\nnext one"
    parts = [text[a:b] for a, b in checks.sentence_spans(text)]
    assert parts == ["About 1.2 million records were copied.", "Patch now! e.g. this stays.", "New line", "next one"]


# ---- 2. "Not in source" ----------------------------------------------------------------------------


def flagged(text: str, sheet, known) -> list[str]:
    content = {"tweets": [{"text": text, "fact_ids": ["F3"]}]}
    quality = checks.check_output("x_thread", content, sheet, known)
    return [f["text"] for f in quality["not_in_source"]]


def test_values_from_the_source_are_not_flagged(sheet, known):
    text = ("As of 28 September 2026, 42 hospitals in five states reported disruption; 1.2 million records at risk; "
            "CVE-2026-XXXXX; 203.0.113.45; recovery in under 2 days; report within 6 hours.")
    assert flagged(text, sheet, known) == []


def test_values_not_in_the_source_are_flagged(sheet, known):
    text = (f"43 hospitals were hit on 23 September; CVE-2026-1234 used 10.0.0.1 and {SHA256[:-1]}0; "
            "fifteen states; 2025.")  # (small numbers like "seven" are in the source: its sections are numbered)
    assert flagged(text, sheet, known) == ["43", "23 September", "CVE-2026-1234", "10.0.0.1", SHA256[:-1] + "0",
                                           "fifteen", "2025"]


def test_a_known_date_with_a_wrong_year_is_flagged(sheet, known):
    assert flagged("First exploited on 22 September 2025.", sheet, known) == ["22 September 2025"]


# ---- 3. consistency across outputs -----------------------------------------------------------------


def checked_outputs(sheet, known, change=None) -> list[dict]:
    outputs = []
    for number, output_type in enumerate(("x_thread", "linkedin_post", "advisory")):
        content = reply(output_type)
        if change and output_type == change[0]:
            change[1](content)
        quality = checks.check_output(output_type, content, sheet, known)
        outputs.append({"id": number + 1, "type": output_type, "label": output_type, "sentences": quality["sentences"]})
    return outputs


def test_mock_outputs_agree(sheet, known):
    result = checks.check_consistency(sheet, checked_outputs(sheet, known))
    assert result["ok"] and result["mismatches"] == []
    top = result["agreed"][0]
    assert (top["fact_id"], top["value"]) == ("F3", "42 hospitals")
    assert set(top["outputs"]) == {"x_thread", "linkedin_post", "advisory"}


def test_a_different_number_for_the_same_fact_is_a_mismatch(sheet, known):
    def say_43(content):
        content["paragraphs"][0]["text"] = "43 hospitals in five states have reported ransomware disruption since 22 September 2026."

    result = checks.check_consistency(sheet, checked_outputs(sheet, known, ("linkedin_post", say_43)))
    assert not result["ok"]
    [mismatch] = result["mismatches"]
    assert mismatch["fact_id"] == "F3"
    assert mismatch["found"] == "43 hospitals"
    assert mismatch["expected"] == ["42 hospitals"]
    assert (mismatch["output_id"], mismatch["output_type"]) == (2, "linkedin_post")
    assert mismatch["sentence_id"] == "s1"


def test_a_number_that_belongs_to_another_fact_is_not_a_mismatch(sheet, known):
    def mention_11(content):
        content["tweets"][0]["text"] = "42 hospitals in five states have reported disruption; 11 hospitals used paper records."

    result = checks.check_consistency(sheet, checked_outputs(sheet, known, ("x_thread", mention_11)))
    assert result["mismatches"] == []


# ---- 4. quality score ------------------------------------------------------------------------------


def test_quality_score_of_a_clean_output(sheet, known):
    quality = checks.check_output("advisory", reply("advisory"), sheet, known)
    assert quality["score"] == 100
    assert {k: (p["points"], p["max"]) for k, p in quality["score_parts"].items()} == {
        "linked": (40, 40), "quotes": (25, 25), "values": (20, 20), "format": (15, 15)}
    assert quality["explanation"].startswith("Quality 100 out of 100.")


def test_quality_score_goes_down_for_each_problem(sheet, known):
    content = {"tweets": [
        {"text": "42 hospitals in five states have reported disruption.", "fact_ids": ["F3"]},
        {"text": "Cyber hygiene matters to everyone.", "fact_ids": []},                     # not linked
        # linked, but a number that is not in the source, and too long for X
        {"text": "43 hospitals in five states have reported disruption" + ", more disruption" * 20 + ".", "fact_ids": ["F3"]},
    ]}
    quality = checks.check_output("x_thread", content, sheet, known)
    parts = quality["score_parts"]
    assert quality["unlinked"] == ["Cyber hygiene matters to everyone."]
    assert parts["values"]["problems"] == 1 and parts["values"]["points"] == 13
    assert parts["format"]["done"] == 2 and parts["format"]["total"] == 3  # post 3 is over 280 characters
    assert quality["score"] == sum(p["points"] for p in parts.values()) < 90
    assert "Post 3 is" in quality["explanation"] and "not linked" not in quality["explanation"]
    assert any("280" in w for w in quality["warnings"])


def test_a_truncated_answer_loses_format_points(sheet, known):
    whole = checks.check_output("x_thread", reply("x_thread"), sheet, known)
    cut = checks.check_output("x_thread", reply("x_thread"), sheet, known, truncated=True)
    assert cut["score"] < whole["score"]


# ---- 5-7. edit, versions, regenerate, downloads (through the API) ------------------------------------


@pytest.fixture()
def job() -> dict:
    created = client.post("/api/jobs", data={"text": SAMPLE, "outputs": ["x_thread", "linkedin_post", "video_package"],
                                             "title": "Trust test"}).json()
    finished = wait_for(created["id"])
    assert finished["status"] == "ready", finished["error"]
    return finished


def output_of(job: dict, output_type: str) -> dict:
    return next(o for o in job["outputs"] if o["type"] == output_type)


def test_job_has_scores_and_consistency(job):
    assert 0 < job["quality_score"] <= 100
    assert job["consistency"]["ok"] is True
    linkedin = output_of(job, "linkedin_post")
    assert linkedin["version"] == 1 and linkedin["origin"] == "ai"
    assert linkedin["quality_score"] == linkedin["quality"]["score"] < 100  # the planted unlinked sentence
    assert {"path": ["tweets", 0, "text"], "label": "Post 1", "text": reply("x_thread")["tweets"][0]["text"]} \
        in output_of(job, "x_thread")["fields"]


def test_source_text_is_served_for_the_trace_panel(job):
    source = client.get(f"/api/jobs/{job['id']}/sources/S1").json()
    fact = job["fact_sheet"]["key_facts"][2]
    page = source["pages"][fact["page"] - 1]
    assert words(page[fact["start"] : fact["end"]]) == words(fact["quote"])
    assert client.get(f"/api/jobs/{job['id']}/sources/S9").status_code == 404


def test_edit_makes_a_new_version_and_runs_the_checks_again(job):
    x_thread = output_of(job, "x_thread")
    url = f"/api/jobs/{job['id']}/outputs/{x_thread['id']}"
    new_text = "Alert: 43 hospitals in five states have reported ransomware disruption."
    response = client.put(url, json={"fields": [{"path": ["tweets", 0, "text"], "text": new_text},
                                                {"path": ["tweets", 3, "text"], "text": ""}]})  # empty = remove
    assert response.status_code == 200, response.text
    edited = output_of(response.json(), "x_thread")

    assert (edited["version"], edited["origin"], edited["origin_label"]) == (2, "human", "Edited by human")
    assert edited["content"]["tweets"][0]["text"] == new_text
    assert len(edited["content"]["tweets"]) == 3
    assert [f["text"] for f in edited["quality"]["not_in_source"]] == ["43"]
    assert edited["quality_score"] < x_thread["quality_score"]
    assert response.json()["consistency"]["mismatches"][0]["found"] == "43 hospitals"

    # the old version is kept and can be viewed
    versions = client.get(f"{url}/versions").json()
    assert [(v["version"], v["origin"]) for v in versions] == [(2, "human"), (1, "ai")]
    old = client.get(f"{url}/versions/1").json()
    assert old["content"]["tweets"][0]["text"] == x_thread["content"]["tweets"][0]["text"]
    assert old["quality"]["score"] == x_thread["quality_score"]

    # downloads use the latest version
    txt = client.get(f"{url}/download", params={"format": "txt"}).text
    assert new_text in txt and "1/3" in txt


def test_editing_narration_updates_the_subtitles(job):
    video = output_of(job, "video_package")
    url = f"/api/jobs/{job['id']}/outputs/{video['id']}"
    new = "Forty-two hospitals were disrupted. Patch your gateways now."
    edited = output_of(client.put(url, json={"fields": [{"path": ["scenes", 0, "narration"], "text": new}]}).json(),
                       "video_package")
    assert [s["text"] for s in edited["content"]["subtitles"][:2]] == ["Forty-two hospitals were disrupted.",
                                                                      "Patch your gateways now."]
    srt = client.get(f"{url}/download", params={"format": "srt"}).text
    assert "Patch your gateways now." in srt


def test_bad_edits_are_refused(job):
    x_thread = output_of(job, "x_thread")
    url = f"/api/jobs/{job['id']}/outputs/{x_thread['id']}"
    assert client.put(url, json={"fields": [{"path": ["tweets", 0, "fact_ids"], "text": "F9"}]}).status_code == 400
    same = x_thread["content"]["tweets"][0]["text"]
    assert client.put(url, json={"fields": [{"path": ["tweets", 0, "text"], "text": same}]}).status_code == 400


def test_regenerate_writes_one_output_again_and_keeps_the_old_version(job):
    linkedin, x_thread = output_of(job, "linkedin_post"), output_of(job, "x_thread")
    url = f"/api/jobs/{job['id']}/outputs/{linkedin['id']}"
    client.put(url, json={"fields": [{"path": ["paragraphs", 3, "text"], "text": ""}]})  # v2: human edit

    response = client.post(f"{url}/regenerate")
    assert response.status_code == 200
    finished = wait_for(job["id"])
    regenerated = output_of(finished, "linkedin_post")
    assert (regenerated["version"], regenerated["origin"]) == (3, "regenerated")
    assert regenerated["content"] == reply("linkedin_post")  # the mock AI writes the same text again
    # the other output was not touched
    assert output_of(finished, "x_thread")["finished_at"] == x_thread["finished_at"]
    assert [v["origin"] for v in client.get(f"{url}/versions").json()] == ["regenerated", "human", "ai"]


# ---- 8. export fixes -----------------------------------------------------------------------------------


def info(title: str) -> ExportInfo:
    return ExportInfo(job_id=142, job_title=title, date="30 Sep 2026", tlp="AMBER", output_type="presentation",
                      output_label="Presentation")


def test_pptx_footer_fits_on_one_line(tmp_path):
    long_title = "Incident report: ransomware campaign affecting hospital networks across five states and more"
    text = footer_text(info(long_title), 1, 6)
    assert text.startswith("Job #142 · Incident report") and text.endswith("… · 30 Sep 2026 · 1/6")
    assert footer_width(text) <= FOOTER_TEXT_POINTS

    short = footer_text(info("Hospital ransomware"), 3, 6)
    assert short == "Job #142 · Hospital ransomware · 30 Sep 2026 · 3/6"

    path = write_pptx(info(long_title), reply("presentation"), tmp_path / "deck.pptx")
    deck = Presentation(str(path))
    for number, slide in enumerate(deck.slides, start=1):
        footer = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.endswith(f"{number}/6"))
        assert footer.text_frame.word_wrap is False
        assert footer_width(footer.text_frame.text) <= FOOTER_TEXT_POINTS


def test_pdf_hash_stays_on_one_line(job):
    created = client.post("/api/jobs", data={"text": SAMPLE, "outputs": ["advisory"]}).json()
    advisory = output_of(wait_for(created["id"]), "advisory")
    pdf = client.get(f"/api/jobs/{created['id']}/outputs/{advisory['id']}/download", params={"format": "pdf"}).content
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert SHA256 in text.splitlines() or any(line.strip().endswith(SHA256) for line in text.splitlines())


def test_old_databases_get_the_new_columns(tmp_path, monkeypatch):
    """A database made before Stage 5 gets the new columns when the app starts; rows are kept."""
    from sqlalchemy import create_engine, inspect, text

    from app import db

    old = create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    with old.begin() as connection:
        connection.execute(text("CREATE TABLE outputs (id INTEGER PRIMARY KEY, job_id INTEGER, type VARCHAR(30))"))
        connection.execute(text("INSERT INTO outputs (id, job_id, type) VALUES (1, 1, 'x_thread')"))
    monkeypatch.setattr(db, "engine", old)
    db.init_db()
    columns = {c["name"] for c in inspect(old).get_columns("outputs")}
    assert {"version", "origin", "quality_score", "quality_json"} <= columns
    with old.connect() as connection:
        assert connection.execute(text("SELECT type, version, origin FROM outputs")).one() == ("x_thread", 0, "ai")
    assert "output_versions" in inspect(old).get_table_names()
