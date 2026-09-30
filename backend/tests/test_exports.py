"""Stage 4 tests: every output becomes a real file that opens, via the download API (mock AI).

Each file is re-opened with the library that reads that format (python-docx, pypdf,
python-pptx, Pillow, zipfile) and checked for the job title, the footer, the TLP label and
the QR code placeholder, and that no fact ids (F1, A2 ...) leaked into the visible text.
"""

import io
import re
import zipfile

import docx
import pytest
from fastapi.testclient import TestClient

from tests.auth_helpers import signed_in_client
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from app.config import settings
from app.db import Job, Output, SessionLocal
from app.exporters import FORMATS
from app.exporters.common import FOOTER, ExportInfo
from app.exporters.infographic import HEIGHT, WIDTH, write_png
from app.main import app
from tests.canned import canned_ai, reply
from tests.helpers import SAMPLE_REPORT
from tests.test_jobs_api import wait_for

client = signed_in_client("operator")


@pytest.fixture(scope="module", autouse=True)
def fixed_answers():
    """These tests check scores, edits and files against known texts, so the AI gives fixed answers
    (tests/canned.py) instead of building them from the source."""
    with canned_ai():
        yield
TITLE = "Hospital ransomware"
FACT_ID = re.compile(r"\b[FA]\d{1,2}\b")  # fact ids must stay in the JSON only


@pytest.fixture(scope="module")
def job() -> dict:
    """One finished job with all 7 outputs, marked TLP:AMBER (normally the safety check sets this)."""
    all_outputs = [o["key"] for o in client.get("/api/options").json()["output_types"]]
    created = client.post("/api/jobs", data={"text": SAMPLE_REPORT.read_text(encoding="utf-8"),
                                             "outputs": all_outputs, "title": TITLE}).json()
    finished = wait_for(created["id"])
    assert finished["status"] == "ready", finished["error"]
    with SessionLocal() as db:
        db.get(Job, finished["id"]).tlp = "AMBER"
        db.commit()
    return client.get(f"/api/jobs/{finished['id']}").json()


def download(job: dict, output_type: str, fmt: str):
    output = next(o for o in job["outputs"] if o["type"] == output_type)
    response = client.get(f"/api/jobs/{job['id']}/outputs/{output['id']}/download", params={"format": fmt})
    assert response.status_code == 200, response.text
    return response


def check_labels(text: str) -> None:
    assert TITLE in text
    assert FOOTER in text
    assert "TLP:AMBER" in text
    assert not FACT_ID.search(text), f"fact id leaked: {FACT_ID.search(text).group()}"


def docx_text(data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    section = document.sections[0]
    parts += [p.text for p in section.header.paragraphs + section.footer.paragraphs]
    return "\n".join(parts)


def test_formats_are_listed_on_each_output(job):
    assert {o["type"]: o["formats"] for o in job["outputs"]} == FORMATS


@pytest.mark.parametrize("output_type", ["advisory", "executive_summary", "video_package"])
def test_word_files_open(job, output_type):
    response = download(job, output_type, "docx")
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    assert f'filename="job{job["id"]}-{output_type.replace("_", "-")}.docx"' in response.headers["content-disposition"]
    text = docx_text(response.content)
    check_labels(text)
    assert "QR code" in text
    if output_type == "advisory":
        assert "Recommended actions" in text and "203.0.113.45" in text  # indicators table
    if output_type == "video_package":
        assert "Storyboard" in text and "On-screen text" in text


@pytest.mark.parametrize("output_type", ["advisory", "executive_summary"])
def test_pdf_files_open(job, output_type):
    response = download(job, output_type, "pdf")
    assert response.headers["content-type"] == "application/pdf"
    reader = PdfReader(io.BytesIO(response.content))
    text = "\n".join(page.extract_text() for page in reader.pages)
    check_labels(text)
    assert "QR code" in text
    fonts = {str(font.get_object()["/BaseFont"]) for page in reader.pages for font in page["/Resources"]["/Font"].values()}
    assert any("Poppins" in f for f in fonts) and any("Hind" in f for f in fonts), fonts  # our fonts are embedded


def test_presentation_opens_with_notes_on_every_slide(job):
    response = download(job, "presentation", "pptx")
    deck = Presentation(io.BytesIO(response.content))
    content = next(o for o in job["outputs"] if o["type"] == "presentation")["content"]
    slides = list(deck.slides)
    assert len(slides) == len(content["slides"]) + 2  # title slide + content slides + closing slide
    for slide in slides:
        assert slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
        text = "\n".join(shape.text_frame.text for shape in slide.shapes if shape.has_text_frame)
        check_labels(text)
    title_slide_text = "\n".join(s.text_frame.text for s in slides[0].shapes if s.has_text_frame)
    assert content["title"] in title_slide_text and "QR code" in title_slide_text
    assert content["slides"][0]["speaker_notes"] == slides[1].notes_slide.notes_text_frame.text
    # the notes master is listed in presentation.xml (without it, macOS Quick Look hangs on the file)
    assert b"notesMasterIdLst" in zipfile.ZipFile(io.BytesIO(response.content)).read("ppt/presentation.xml")


def test_infographic_png(job):
    response = download(job, "infographic", "png")
    assert response.headers["content-type"] == "image/png"
    image = Image.open(io.BytesIO(response.content))
    assert image.format == "PNG" and image.size == (WIDTH, HEIGHT) == (1080, 1350)
    image.verify()


def test_infographic_preview_is_inline(job):
    output = next(o for o in job["outputs"] if o["type"] == "infographic")
    response = client.get(f"/api/jobs/{job['id']}/outputs/{output['id']}/download?format=png&inline=true")
    assert response.headers["content-disposition"].startswith("inline")


@pytest.mark.parametrize("layout", ["number_grid", "vertical_steps", "timeline", "something_else"])
def test_every_infographic_layout_draws(tmp_path, layout):
    content = reply("infographic") | {"layout": layout}
    # a very long headline and many steps must still fit (the drawing shrinks)
    content["headline"] = "A very long headline that goes on and on to test how the image copes with it"
    content["steps"] = content["steps"] + [{"text": "One more step with quite a lot of words in it to wrap", "fact_ids": []}]
    info = ExportInfo(1, "Test job", "30 Sep 2026", None, "infographic", "Infographic")
    path = write_png(info, content, tmp_path / "info.png")
    assert Image.open(path).size == (1080, 1350)


def test_subtitles_srt(job):
    response = download(job, "video_package", "srt")
    text = response.content.decode("utf-8")
    cues = [cue.split("\n") for cue in text.strip().split("\n\n")]
    assert [int(cue[0]) for cue in cues] == list(range(1, len(cues) + 1))
    assert all(re.fullmatch(r"\d\d:\d\d:\d\d,\d{3} --> \d\d:\d\d:\d\d,\d{3}", cue[1]) for cue in cues)
    assert cues[0][1].startswith("00:00:00,000")
    check_labels(cues[-1][2] + "\n" + cues[-1][3])  # the closing cue carries the labels
    assert FOOTER not in "\n".join(line for cue in cues[:-1] for line in cue)


def test_x_thread_txt_is_numbered(job):
    text = download(job, "x_thread", "txt").content.decode("utf-8")
    check_labels(text)
    tweets = next(o for o in job["outputs"] if o["type"] == "x_thread")["content"]["tweets"]
    for number in range(1, len(tweets) + 1):
        assert f"{number}/{len(tweets)} " in text
    assert "QR code" in text


def test_linkedin_txt(job):
    text = download(job, "linkedin_post", "txt").content.decode("utf-8")
    check_labels(text)
    assert "#CyberSecurity" in text


def test_campaign_kit_zip(job):
    response = client.get(f"/api/jobs/{job['id']}/kit.zip")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/zip"
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert archive.testzip() is None  # every file in the zip is intact
    expected = {f"job{job['id']}-{t.replace('_', '-')}.{fmt}" for t, formats in FORMATS.items() for fmt in formats}
    assert set(archive.namelist()) == expected | {"README.txt"}
    readme = archive.read("README.txt").decode("utf-8")
    check_labels(readme)
    assert len(re.findall(r"\b[0-9a-f]{64}\b", readme)) == len(expected)  # a SHA-256 per file
    # the kit's files re-open
    Presentation(io.BytesIO(archive.read(f"job{job['id']}-presentation.pptx")))
    docx.Document(io.BytesIO(archive.read(f"job{job['id']}-advisory.docx")))


def test_files_are_saved_under_data_exports(job):
    folder = settings.data_dir / "jobs" / str(job["id"]) / "exports"
    client.get(f"/api/jobs/{job['id']}/kit.zip")
    names = {p.name for p in folder.iterdir()}
    assert f"job{job['id']}-campaign-kit.zip" in names and f"job{job['id']}-advisory.pdf" in names
    assert not any(name.endswith(".part") for name in names)  # no half-written files left


def test_download_errors(job):
    advisory = next(o for o in job["outputs"] if o["type"] == "advisory")
    base = f"/api/jobs/{job['id']}/outputs/{advisory['id']}/download"
    assert client.get(base, params={"format": "pptx"}).status_code == 400  # not a format of an advisory
    assert "pdf, docx" in client.get(base, params={"format": "pptx"}).json()["detail"]
    assert client.get(base).status_code == 422  # format is required
    assert client.get(f"/api/jobs/999999/outputs/{advisory['id']}/download?format=pdf").status_code == 404
    assert client.get("/api/jobs/999999/kit.zip").status_code == 404


def test_unfinished_output_cannot_be_downloaded():
    with SessionLocal() as db:
        job = Job(title="Not finished", status="generating")
        db.add(job)
        db.flush()
        output = Output(job=job, type="advisory", status="queued")
        db.add(output)
        db.commit()
        job_id, output_id = job.id, output.id
    response = client.get(f"/api/jobs/{job_id}/outputs/{output_id}/download?format=pdf")
    assert response.status_code == 409
    assert "not finished" in response.json()["detail"]
    assert client.get(f"/api/jobs/{job_id}/kit.zip").status_code == 409
