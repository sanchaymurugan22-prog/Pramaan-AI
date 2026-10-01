"""Stage 9B part 5b: the letterhead (office name and logo) on every exported file and as "Issued by", logo
checks, the public page information, and backups (encrypted database copy that opens with the key)."""

import io
import uuid
import zipfile

import docx
import pypdf
import sqlcipher3
from PIL import Image
from pptx import Presentation

from app import app_settings, branding, crypto
from tests.auth_helpers import signed_in_client
from tests.helpers import SAMPLE_REPORT
from tests.test_jobs_api import wait_for

admin = signed_in_client("admin", "files.admin")
operator = signed_in_client("operator", "files.operator")
reviewer = signed_in_client("reviewer", "files.reviewer")


def png(width=200, height=80, colour=(30, 47, 143)) -> bytes:
    out = io.BytesIO()
    Image.new("RGB", (width, height), colour).save(out, "PNG")
    return out.getvalue()


def finished_job(outputs: list[str]) -> dict:
    job = operator.post("/api/jobs", data={"text": SAMPLE_REPORT.read_text(encoding="utf-8") + uuid.uuid4().hex,
                                           "outputs": outputs}).json()
    return wait_for(job["id"])


def test_logo_checks():
    assert admin.post("/api/admin/letterhead/logo", files={"file": ("x.png", b"not an image", "image/png")}).status_code == 400
    assert admin.post("/api/admin/letterhead/logo", files={"file": ("tiny.png", png(10, 10), "image/png")}).status_code == 400
    gif = io.BytesIO()
    Image.new("RGB", (64, 64)).save(gif, "GIF")
    assert admin.post("/api/admin/letterhead/logo", files={"file": ("a.gif", gif.getvalue(), "image/gif")}).status_code == 400


def test_letterhead_on_every_file_and_as_issuer():
    try:
        assert admin.put("/api/admin/letterhead", json={"office_name": "CERT-Test Regional Office"}).json()["effective_name"] == \
            "CERT-Test Regional Office"
        uploaded = admin.post("/api/admin/letterhead/logo", files={"file": ("logo.png", png(1200, 400), "image/png")})
        assert uploaded.status_code == 200 and uploaded.json()["has_logo"]
        stored = branding.LOGO_PATH.read_bytes()
        assert not stored.startswith(b"\x89PNG")  # stored encrypted
        assert Image.open(io.BytesIO(crypto.read_file(branding.LOGO_PATH))).width == 600  # made smaller
        assert operator.get("/api/letterhead/logo.png").headers["content-type"] == "image/png"

        job = finished_job(["advisory", "presentation", "linkedin_post", "infographic"])
        files = {o["type"]: o["id"] for o in job["outputs"]}
        base = f"/api/jobs/{job['id']}/outputs"
        pdf = pypdf.PdfReader(io.BytesIO(operator.get(f"{base}/{files['advisory']}/download?format=pdf").content))
        assert "CERT-Test Regional Office · Advisory" in pdf.pages[0].extract_text()
        assert pdf.pages[0].images  # the logo is drawn
        word = docx.Document(io.BytesIO(operator.get(f"{base}/{files['advisory']}/download?format=docx").content))
        header = " ".join(p.text for p in word.sections[0].header.paragraphs)
        assert "CERT-Test Regional Office · Advisory" in header
        deck = Presentation(io.BytesIO(operator.get(f"{base}/{files['presentation']}/download?format=pptx").content))
        first = deck.slides[0]
        assert any(s.has_text_frame and s.text_frame.text == "CERT-Test Regional Office" for s in first.shapes)
        assert any(s.shape_type == 13 for s in first.shapes)  # 13 = picture
        poster = Image.open(io.BytesIO(operator.get(f"{base}/{files['infographic']}/download?format=png").content))
        assert poster.getpixel((90, 43)) == (30, 47, 143)  # inside the logo, top left
        text = operator.get(f"{base}/{files['linkedin_post']}/download?format=txt").text
        assert text.startswith("CERT-Test Regional Office · LinkedIn post")

        # signing says who issued it
        operator.post(f"/api/jobs/{job['id']}/submit", json={})
        record = reviewer.post(f"/api/jobs/{job['id']}/review", json={"decision": "approve"}).json()["record"]
        assert reviewer.get(f"/api/records/{record['record_no']}").json()["issuing_office"] == "CERT-Test Regional Office"
    finally:
        admin.delete("/api/admin/letterhead/logo")
        app_settings.put("letterhead", app_settings.DEFAULTS["letterhead"])
    assert admin.get("/api/admin/letterhead").json()["has_logo"] is False
    assert operator.get("/api/letterhead/logo.png").status_code == 404


def test_public_page_information():
    body = admin.get("/api/admin/public-page").json()
    assert body["address"].startswith("http") and "issued" in body["records"]
    assert admin.get("/api/admin/records/verify-bundle.zip").status_code == 200
    assert admin.get("/api/admin/public-page").json()["last_export"]


def test_backup_holds_an_encrypted_database_that_opens_with_the_key():
    finished_job(["x_thread"])
    made = admin.post("/api/admin/backups")
    assert made.status_code == 201, made.text
    name = made.json()["made"]
    assert name in [b["name"] for b in admin.get("/api/admin/backups").json()["backups"]]
    download = admin.get(f"/api/admin/backups/{name}")
    archive = zipfile.ZipFile(io.BytesIO(download.content))
    names = archive.namelist()
    assert "pramaan.db" in names and "README.txt" in names and any(n.startswith("jobs/") for n in names)
    assert not any(n.startswith(("backups/", "watch/")) for n in names)
    database = archive.read("pramaan.db")
    assert not database.startswith(b"SQLite format 3")  # encrypted
    path = branding.LOGO_PATH.parent.parent / "restore-test.db"
    path.write_bytes(database)
    try:
        copy = sqlcipher3.connect(str(path))
        copy.execute(f"PRAGMA key = \"x'{crypto.database_key_hex()}'\"")
        assert copy.execute("SELECT count(*) FROM jobs").fetchone()[0] >= 1
        copy.close()
    finally:
        path.unlink()
    # only backups from the list can be downloaded
    assert admin.get("/api/admin/backups/..%2Fpramaan.db").status_code == 404
    assert admin.get("/api/admin/backups/pramaan.db").status_code == 404
